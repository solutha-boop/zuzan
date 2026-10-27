import io, types, asyncio, datetime as dt
import pytest
import openpyxl
from fastapi import UploadFile, HTTPException
from tests.helpers import *
from database import Invoice, InvoiceStatus, Employee, Account, JournalLine, JournalEntry
import csv_import as CI


def _xlsx(sheets):
    """sheets: list of (title, rows); returns xlsx bytes. Cells keep their Python types (numbers stay numbers)."""
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    for title, rows in sheets:
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def _up(b): return UploadFile(filename="x.xlsx", file=io.BytesIO(b))


def _cu(name):
    co = company(name, "general"); return co, types.SimpleNamespace(company_id=co.id, id=1, role="owner")


def test_reader_numbers_dates_and_title_rows():
    b = _xlsx([("Sheet1", [["Solutha Bookkeeping - Export"], [], ["Code", "Amount", "Date", "Flag"],
                           [1001, 1234.5, dt.datetime(2026, 3, 1), True], [2002.0, 10000.0, dt.datetime(2026, 3, 2, 14, 30), False]])])
    h, rows = CI._read_csv(b)
    assert h == ["Code", "Amount", "Date", "Flag"]
    assert rows[0][:2] == ["1001", "1234.5"] and rows[1][:2] == ["2002", "10000"]      # no '2002.0'
    assert CI._parse_date(rows[0][2]) == dt.datetime(2026, 3, 1) and CI._parse_date(rows[1][2]).day == 2


def test_reader_percent_cells_old_xls_empty_and_blank_cover_sheet():
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["Name", "Pct"]); ws.append(["a", 0.075]); ws["B2"].number_format = "0.0%"
    buf = io.BytesIO(); wb.save(buf)
    assert CI._read_csv(buf.getvalue())[1][0][1] == "7.5"                               # shown as 7.5%, stored 0.075
    with pytest.raises(HTTPException) as e:
        CI._read_csv(b"\xd0\xcf\x11\xe0" + b"0" * 100)
    assert e.value.status_code == 400 and ".xlsx" in e.value.detail
    with pytest.raises(HTTPException) as e:
        CI._read_csv(_xlsx([("S", [[None, None]])]))
    assert e.value.status_code == 400
    h, rows = CI._read_csv(_xlsx([("Cover", []), ("Data", [["Name", "Qty"], ["x", 3]])]))   # empty first tab
    assert h == ["Name", "Qty"] and rows == [["x", 3 and "3"]]


def test_corrupt_or_protected_zip_is_a_clean_400():
    with pytest.raises(HTTPException) as e:
        CI._read_csv(b"PK\x03\x04" + b"garbage" * 50)
    assert e.value.status_code == 400


def test_excel_invoice_import_with_real_cell_types():
    co, cu = _cu("xlInv")
    b = _xlsx([("Invoices", [["Invoice Number", "Client Name", "Date", "Total", "VAT", "Status"],
                             ["X-1", "Acme", dt.datetime(2026, 3, 1), 1150, 150, "Paid"],
                             ["X-2", "Beta", dt.datetime(2026, 3, 5), 230.0, 30.0, "Sent"]])])
    r = asyncio.run(CI.import_invoices(file=_up(b), db=db, cu=cu))
    assert r["imported"] == 2 and not r["errors"]
    a = db.query(Invoice).filter_by(company_id=co.id, invoice_number="X-1").one()
    assert a.total_amount == 1150 and a.amount == 1000 and a.status == InvoiceStatus.paid and a.issue_date.month == 3
    assert unbal(co) == 0
    r2 = asyncio.run(CI.import_invoices(file=_up(b), db=db, cu=cu))
    assert r2["imported"] == 0                                                          # not doubled


def test_excel_employee_numbers_stored_as_numbers_do_not_duplicate_staff():
    co, cu = _cu("xlEmp")
    first = _xlsx([("Staff", [["Employee Number", "ID Number", "First Name", "Last Name", "Gross Salary", "Pension Employee Pct"],
                              [1001, 8001015009087, "Ann", "Able", 12345.67, 0.075]])])
    wb = openpyxl.load_workbook(io.BytesIO(first)); wb.active["F2"].number_format = "0.0%"
    buf = io.BytesIO(); wb.save(buf)
    r = asyncio.run(CI.import_employees(file=_up(buf.getvalue()), db=db, cu=cu))
    assert r["created"] == 1 and not r["errors"] and not r["warnings"]
    e = db.query(Employee).filter_by(company_id=co.id).one()
    assert e.employee_number == "1001" and e.id_number == "8001015009087" and e.pension_fund_employee_pct == 7.5
    # next month's file: same people, a raise - must update, not create a second 'Ann'
    second = _xlsx([("Staff", [["Employee Number", "First Name", "Last Name", "Gross Salary"], [1001, "Ann", "Able", 13000]])])
    r = asyncio.run(CI.import_employees(file=_up(second), db=db, cu=cu))
    assert r["updated"] == 1 and r["created"] == 0
    assert db.query(Employee).filter_by(company_id=co.id).count() == 1 and db.query(Employee).filter_by(company_id=co.id).one().gross_salary == 13000


def test_excel_trial_balance_and_payroll_adjustments():
    co, cu = _cu("xlTB")
    tb = _xlsx([("TB", [["Trial Balance as at 28 Feb 2025"], ["Code", "Account Name", "Type", "Debit", "Credit"],
                        [1000, "Bank", "Asset", 8000, None], [2000, "Creditors", "Liability", None, 3000], [3000, "Capital", "Equity", None, 5000]])])
    r = asyncio.run(CI.import_trial_balance(file=_up(tb), as_at="2025-02-28", replace=False, db=db, cu=cu))
    assert r["imported"] == 3 and r["balancing_to_opening_equity"] == 0 and unbal(co) == 0
    a = db.query(Account).filter_by(company_id=co.id, code="1000").one()
    assert sum((l.debit or 0) - (l.credit or 0) for l in db.query(JournalLine).filter_by(account_id=a.id)) == 8000
    asyncio.run(CI.import_employees(file=_up(_xlsx([("S", [["Employee Number", "First Name", "Last Name", "Gross Salary"], [7, "Bo", "Bee", 10000]])])), db=db, cu=cu))
    r = asyncio.run(CI.import_payroll_adjustments(file=_up(_xlsx([("A", [["Employee Number", "New Salary"], [7, 11000.5]])])), db=db, cu=cu))
    assert r["imported"] == 1 and db.query(Employee).filter_by(company_id=co.id).one().gross_salary == 11000.5
